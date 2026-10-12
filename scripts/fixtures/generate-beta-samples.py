from pathlib import Path
import json, hashlib, sys, datetime
from zipfile import ZipFile
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.pagesizes import A4
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "public" / "autofill" / "samples"
PRIVATE = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
PRIVATE.mkdir(parents=True, exist_ok=True)
NAME = "沈星澄-Beta增强虚构简历"
TITLE = "虚构 Beta 增强简历"
NOTICE = "全部姓名、联系方式、院校、机构、证书和经历均为原创虚构内容，仅供本机验收，不用于真实投递。"

contact = {
    "name": "沈星澄", "givenName": "Xingcheng", "familyName": "Shen",
    "email": "shen.xingcheng@example.com", "phone": "19900003210",
    "gender": "女", "politicalStatus": "群众", "birthDate": "1998-11-08",
    "province": "浙江", "city": "杭州", "district": "西湖",
    "website": "https://example.com/resume-beta",
}
education = [
    {"school": "虚构青禾理工大学", "degree": "本科", "academicDegree": "工学学士", "major": "软件工程", "startDate": "2017-09", "endDate": "2021-06"},
    {"school": "虚构望舒研究学院", "degree": "硕士", "academicDegree": "工学硕士", "major": "计算机技术", "startDate": "2021-09", "endDate": "2024-06"},
]
internships = [
    {"organization": "虚构星桥数据工作室", "title": "数据整理实习生", "startDate": "2020-07", "endDate": "2020-09", "description": "仅为测试：整理 CSV 资料，核对字段映射并记录缺失项。"},
    {"organization": "虚构林间软件实验室", "title": "开发实习生", "startDate": "2023-02", "endDate": "2023-05", "description": "仅为测试：维护本机表单，补充输入校验并复查保存结果。"},
]
work = [
    {"organization": "虚构清流软件公司", "title": "软件工程师", "startDate": "2024-07", "endDate": "2025-06", "description": "仅为测试：维护资料工具，补充导入预览并核对来源记录。"},
    {"organization": "虚构远帆技术公司", "title": "应用开发工程师", "startDate": "2025-07", "endDate": "2026-09", "description": "仅为测试：优化本机编辑流程，检查加密备份的恢复内容。"},
]
projects = [
    {"name": "虚构资料核对台", "role": "开发成员", "startDate": "2022-03", "endDate": "2022-06", "description": "仅为测试：设计字段核对页面，保留原文并展示待确认状态。"},
    {"name": "虚构学习记录册", "role": "界面开发", "startDate": "2023-07", "endDate": "2023-10", "description": "仅为测试：制作记录卡片，验证修改后的保存与重开流程。"},
]
skills = [{"name": "JavaScript", "level": "熟练"}, {"name": "SQL", "level": "熟悉"}]
certificates = [
    {"name": "虚构软件测试结业证", "issuer": "虚构青禾培训中心", "date": "2024-05-12"},
    {"name": "虚构数据整理结业证", "issuer": "虚构星桥课堂", "date": "2025-03-20"},
]
languages = [
    {"name": "英语", "level": "熟练", "certificate": "仅为测试的 CET-6 成绩 540"},
    {"name": "日语", "level": "基础", "certificate": "仅为测试的 JLPT N3"},
]
about = {
    "introduction": "仅为测试：我参与资料工具维护，习惯先核对来源，再整理表述与保存结果。",
    "other": "仅为测试：可提供虚构作品演示，所有资料只用于本机功能检查。",
}
expected = {"contact": contact, "education": education, "internships": internships, "work": work, "projects": projects, "skills": skills, "certificates": certificates, "languages": languages, "about": about}
contact_labels = {"name": "姓名", "givenName": "英文名", "familyName": "英文姓", "email": "电子邮箱", "phone": "手机号码", "gender": "性别", "politicalStatus": "政治面貌", "birthDate": "出生日期", "province": "现居省份", "city": "居住城市", "district": "所在区县", "website": "作品集网址"}
labels = {
    "education": {"school": "院校名称", "degree": "学历", "academicDegree": "学位", "major": "专业", "startDate": "入学时间", "endDate": "毕业时间"},
    "internships": {"organization": "实习单位", "title": "实习职位", "startDate": "开始时间", "endDate": "结束时间", "description": "经历描述"},
    "work": {"organization": "公司名称", "title": "职位名称", "startDate": "开始时间", "endDate": "结束时间", "description": "工作描述"},
    "projects": {"name": "项目名称", "role": "担任角色", "startDate": "开始时间", "endDate": "结束时间", "description": "项目描述"},
    "skills": {"name": "技能名称", "level": "熟练程度"},
    "certificates": {"name": "证书名称", "issuer": "颁发机构", "date": "获得时间"},
    "languages": {"name": "语种", "level": "熟练程度", "certificate": "语言证书"},
}
section_names = {"education": "教育经历", "internships": "实习经历", "work": "工作经历", "projects": "项目经历", "skills": "技能", "certificates": "证书", "languages": "语言能力"}
pages = [[("title", TITLE), ("notice", NOTICE), ("heading", "个人资料")]]
pages[0] += [("body", contact_labels[k] + "：" + v) for k, v in contact.items()]
def section(key):
    result = [("heading", section_names[key])]
    for i, row in enumerate(expected[key]):
        result += [("body", labels[key][k] + "：" + v) for k, v in row.items()]
        if i < len(expected[key]) - 1:
            result.append(("space", ""))
    return result
pages[0] += section("education")
pages.append(section("internships") + section("work"))
pages.append(section("projects") + section("skills") + section("certificates") + section("languages") + [("heading", "自我介绍"), ("body", about["introduction"]), ("heading", "其他补充"), ("body", about["other"])])

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.top_margin = s.bottom_margin = Mm(19)
s.left_margin = s.right_margin = Mm(21)
for name in ["Normal", "Title", "Heading 1"]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(2)
doc.styles["Normal"].paragraph_format.line_spacing = 1.08
doc.styles["Title"].font.size = Pt(18)
doc.styles["Title"].paragraph_format.space_after = Pt(7)
doc.styles["Heading 1"].font.size = Pt(12)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(3)
doc.styles["Heading 1"].paragraph_format.keep_with_next = True
for pi, lines in enumerate(pages):
    if pi:
        doc.add_page_break()
    for kind, text in lines:
        if kind == "space":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 0.25
            continue
        p = doc.add_paragraph(text, "Title" if kind == "title" else "Heading 1" if kind == "heading" else "Normal")
        if kind == "notice":
            p.runs[0].font.size = Pt(9)
        if kind == "heading":
            p.paragraph_format.keep_with_next = True
doc.core_properties.title = TITLE
doc.core_properties.subject = "原创虚构 Beta 增强字段本机导入验收"
doc.core_properties.author = "职有引力虚构验收样本"
doc.core_properties.last_modified_by = "职有引力虚构验收样本"
doc.core_properties.created = datetime.datetime(2026, 10, 12, tzinfo=datetime.timezone.utc)
doc.core_properties.modified = datetime.datetime(2026, 10, 12, tzinfo=datetime.timezone.utc)
docx = OUT / (NAME + ".docx")
doc.save(docx)

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
pdf = OUT / (NAME + ".pdf")
canvas = Canvas(str(pdf), pagesize=A4, pageCompression=1, invariant=1)
canvas.setTitle(TITLE)
canvas.setAuthor("职有引力虚构验收样本")
for lines in pages:
    y = 794
    for kind, text in lines:
        if kind == "space":
            y -= 5
            continue
        size = 18 if kind == "title" else 12 if kind == "heading" else 8.3 if kind == "notice" else 10.5
        if kind == "heading":
            y -= 7
        canvas.setFont("STSong-Light", size)
        width = pdfmetrics.stringWidth(text, "STSong-Light", size)
        if width > 480:
            raise RuntimeError(f"Line too wide ({width:.1f} pt): {text}")
        if y < 52:
            raise RuntimeError("Page content exceeds bottom margin")
        canvas.drawString(58, y, text)
        y -= 25 if kind == "title" else 19 if kind == "heading" else 14.5
    canvas.showPage()
canvas.save()

with ZipFile(docx) as z:
    embedded_fonts = [n for n in z.namelist() if n.startswith("word/fonts/")]
    if embedded_fonts:
        raise RuntimeError("Unexpected embedded fonts")
pdf_reader = PdfReader(pdf)
assert len(pdf_reader.pages) == 3
fonts = []
for page in pdf_reader.pages:
    for ref in page["/Resources"]["/Font"].values():
        obj = ref.get_object()
        fonts.append(str(obj.get("/BaseFont", "")))
        for descendant in obj.get("/DescendantFonts", []):
            descriptor = descendant.get_object().get("/FontDescriptor")
            if descriptor:
                descriptor = descriptor.get_object()
                assert not any(k in descriptor for k in ["/FontFile", "/FontFile2", "/FontFile3"])
text = "\n".join(text for lines in pages for kind, text in lines if kind != "space")
(PRIVATE / "expected.private.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(PRIVATE / "original.private.txt").write_text(text + "\n", encoding="utf-8")
info = {
    "createdAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
    "originalContent": True, "fictionalOnly": True, "generator": str(Path(__file__).resolve()),
    "python": sys.executable, "fonts": {"docx": "Microsoft YaHei reference only; no font binaries", "pdf": "standard CID STSong-Light; no embedded Windows font binary"},
    "pdfPages": 3, "pdfFontObjects": sorted(set(fonts)), "docxEmbeddedFonts": embedded_fonts,
    "files": [{"path": str(path), "size": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in [docx, pdf]],
}
(PRIVATE / "provenance.private.json").write_text(json.dumps(info, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(info, ensure_ascii=False, indent=2))

